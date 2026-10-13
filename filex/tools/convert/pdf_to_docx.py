"""``pdf.to-docx``: PDF to an editable Word document, fast mode (spec 6, plan task 32).

PyMuPDF reads each page's text blocks (with font size, bold, italic, colour), its images and its
ruled tables; python-docx writes them as paragraphs, pictures and Word tables, one PDF page per
Word page. The engine choice is measured in docs/benchmarks/pdf-to-docx.md (pdf2docx, the
spec's first pick, is unmaintained and pulls OpenCV and NumPy into the core image).

What it keeps: the text in reading order (top to bottom), paragraphs, short lines as line
breaks, bigger text as headings, run styles, images at their size, ruled tables as tables,
page breaks, and each page's size and orientation (a new section when it changes). What it
drops: exact positions, columns (text side by side is flattened into one column and warned
about as ``LAYOUT_SIMPLIFIED``), vector drawings, form fields, links, fonts.

Untrusted input: python-docx writes plain OOXML from its own template (no macros, no external
relationships); text loses the control characters XML cannot hold. Images larger than
``FILEX_MAX_PIXELS`` are left out (``IMAGES_SKIPPED``), and so are soft masks of another size
than their image; both are checked before conversion to PNG (:func:`_image_bytes`).
"""

from __future__ import annotations

import io
import re
from collections import Counter
from pathlib import Path
from typing import Any

from filex.registry import Edge, Example, Files, StepContext, ToolOptions, tool
from filex.tools._pdf import label, open_mupdf

MARGIN_PT = 54  # 0.75 in
MIN_IMAGE_PT = 4  # smaller placements are spacers or rules
CLIP_DPI = 150  # inline images (no xref) are rendered from the page at this resolution
_XML_BAD = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff￾￿]")


@tool(
    id="pdf.to-docx",
    title="PDF to Word",
    summary="Turn a PDF into an editable Word document. The layout is simplified.",
    category="convert",
    group="core",
    converts=[Edge("pdf", "docx", lossy=True)],
    files=Files(min=1, max=20, each=True),
    options=ToolOptions,
    engines=["pymupdf", "python-docx"],
    examples=[Example(inputs=["pdf/text_10p.pdf"], target="docx")],
    warnings=["NO_TEXT", "LAYOUT_SIMPLIFIED", "IMAGES_SKIPPED", "INPUT_REPAIRED"],
    errors=["INVALID_PDF", "ENCRYPTED_PDF", "TOO_MANY_PAGES"],
    keywords=["word", "doc", "docx", "editable", "office", "msword"],
)
def run(ctx: StepContext, src: Path, opts: ToolOptions) -> list[Path]:
    import docx
    import pymupdf

    what = label(ctx)
    pdf = open_mupdf(ctx, src, what)
    try:
        pages = []
        for i, page in enumerate(pdf):
            ctx.progress(i / len(pdf) * 0.6, f"Reading page {i + 1}")
            pages.append(_read(page, pymupdf))
        body = _body_size(pages)
        doc = docx.Document()
        doc.core_properties.author = doc.core_properties.last_modified_by = ""
        images: dict[int, bytes | None] = {}
        skipped = columns = 0
        for i, (page, data) in enumerate(zip(pdf, pages, strict=True)):
            ctx.progress(0.6 + i / len(pdf) * 0.4, f"Writing page {i + 1}")
            _new_page(doc, page.rect, first=i == 0)
            columns += data["columns"]
            for _y, kind, item in data["items"]:
                if kind == "text":
                    _paragraph(doc, item, body)
                elif kind == "table":
                    _table(doc, item)
                elif not _picture(ctx, doc, pdf, page, item, images, pymupdf):
                    skipped += 1
    finally:
        pdf.close()
    if not any(d["chars"] for d in pages):
        ctx.warn(
            "NO_TEXT",
            f"{what} has no text layer; it looks like a scan. The Word file holds only the page "
            "images. Run OCR on it first, then convert it to Word.",
        )
    if columns:
        ctx.warn(
            "LAYOUT_SIMPLIFIED",
            f"{what} has text in columns on {_n(columns, 'page')}; it is one column in the Word "
            "file.",
        )
    if skipped:
        ctx.warn(
            "IMAGES_SKIPPED",
            f"{_n(skipped, 'image')} of {what} could not be read or exceeded "
            f"{ctx.max_pixels / 1e6:g} megapixels and are left out.",
        )
    target = ctx.out_dir / "document.docx"
    doc.save(target)
    return [target]


def _read(page: Any, pymupdf: Any) -> dict:
    """A page's content as ``(top, kind, item)`` in reading order, plus its character count and
    whether it has text side by side (columns)."""
    tables = [t for t in page.find_tables().tables if t.row_count and t.col_count]
    boxes = [pymupdf.Rect(t.bbox) for t in tables]
    flags = pymupdf.TEXTFLAGS_TEXT | pymupdf.TEXT_DEHYPHENATE
    blocks = []
    for b in page.get_text("dict", flags=flags)["blocks"]:
        r = pymupdf.Rect(b["bbox"])
        centre = pymupdf.Point((r.x0 + r.x1) / 2, (r.y0 + r.y1) / 2)
        if b.get("lines") and not any(centre in box for box in boxes):
            blocks.append(b)
    # Text in MuPDF's order (the content stream's, which keeps columns apart); each table and
    # image goes before the first text block that starts below its top.
    others = [(t.bbox[1], "table", t.extract()) for t in tables]
    for info in page.get_image_info(xrefs=True):
        r = pymupdf.Rect(info["bbox"]) & page.rect
        if r.width >= MIN_IMAGE_PT and r.height >= MIN_IMAGE_PT:
            others.append((r.y0, "image", info | {"bbox": r}))
    others.sort(key=lambda it: it[0])
    items = []
    for b in blocks:
        while others and others[0][0] <= b["bbox"][1]:
            items.append(others.pop(0))
        items.append((b["bbox"][1], "text", b))
    items += others
    chars = sum(len(s["text"].strip()) for b in blocks for ln in b["lines"] for s in ln["spans"])
    return {
        "items": items,
        "chars": chars,
        "columns": _has_columns([pymupdf.Rect(b["bbox"]) for b in blocks]),
    }


def _has_columns(rects: list[Any]) -> bool:
    """Two text blocks of several lines' height next to each other (overlapping rows, apart
    across)."""
    tall = [r for r in rects if r.height > 30]
    for i, a in enumerate(tall):
        for b in tall[i + 1 :]:
            overlap = min(a.y1, b.y1) - max(a.y0, b.y0)
            if overlap > min(a.height, b.height) / 2 and (a.x1 <= b.x0 or b.x1 <= a.x0):
                return True
    return False


def _body_size(pages: list[dict]) -> float:
    """The most common font size, weighted by characters: the size of body text."""
    sizes: Counter[float] = Counter()
    for data in pages:
        for _y, kind, b in data["items"]:
            if kind == "text":
                for ln in b["lines"]:
                    for s in ln["spans"]:
                        sizes[round(s["size"] * 2) / 2] += len(s["text"])
    return sizes.most_common(1)[0][0] if sizes else 11.0


def _new_page(doc: Any, rect: Any, *, first: bool) -> None:
    """Start a Word page for a PDF page: a page break, or a new section when the size changes."""
    from docx.enum.section import WD_ORIENT, WD_SECTION
    from docx.shared import Pt

    section = doc.sections[-1]
    width, height = Pt(rect.width), Pt(rect.height)
    if not first:
        if abs(section.page_width - width) < Pt(1) and abs(section.page_height - height) < Pt(1):
            doc.add_page_break()
            return
        section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.page_width, section.page_height = width, height
    section.orientation = WD_ORIENT.LANDSCAPE if rect.width > rect.height else WD_ORIENT.PORTRAIT
    margin = Pt(min(MARGIN_PT, rect.width / 8, rect.height / 8))
    section.left_margin = section.right_margin = margin
    section.top_margin = section.bottom_margin = margin


def _paragraph(doc: Any, block: dict, body: float) -> None:
    """One text block as one paragraph: lines joined, a short line ends with a line break,
    large text is a heading."""
    from docx.shared import Pt, RGBColor

    spans = [s for ln in block["lines"] for s in ln["spans"] if s["text"].strip()]
    if not spans:
        return
    size = min(s["size"] for s in spans)
    para = doc.add_paragraph()
    if size >= body * 1.25 and len(block["lines"]) <= 3:
        para.style = doc.styles["Heading 1" if size >= body * 1.6 else "Heading 2"]
    x1 = block["bbox"][2]
    width = max(x1 - block["bbox"][0], 1)
    for n, line in enumerate(block["lines"]):
        if n:
            prev = block["lines"][n - 1]
            if x1 - prev["bbox"][2] > width / 4:
                para.add_run().add_break()
            elif not prev["spans"] or not prev["spans"][-1]["text"].endswith((" ", "-")):
                para.add_run(" ")
        for s in line["spans"]:
            text = _XML_BAD.sub("", s["text"].replace("\r", " ").replace("\n", " "))
            if not text:
                continue
            run = para.add_run(text)
            run.font.size = Pt(round(s["size"] * 2) / 2)
            run.font.bold = bool(s["flags"] & 16) or None
            run.font.italic = bool(s["flags"] & 2) or None
            run.font.superscript = bool(s["flags"] & 1) or None
            run.font.color.rgb = RGBColor.from_string(f"{s['color'] & 0xFFFFFF:06X}")


def _table(doc: Any, rows: list[list[str | None]]) -> None:
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = doc.styles["Table Grid"]
    for row, cells in zip(table.rows, rows, strict=True):
        for cell, text in zip(row.cells, cells, strict=False):
            cell.text = _XML_BAD.sub("", text or "")


def _picture(
    ctx: StepContext,
    doc: Any,
    pdf: Any,
    page: Any,
    info: dict,
    cache: dict[int, bytes | None],
    pymupdf: Any,
) -> bool:
    """Add one image placement at its size on the page (at most the text width). False when
    it was left out (too many pixels, unreadable)."""
    from docx.shared import Pt

    xref, box = info["xref"], info["bbox"]
    if info["width"] * info["height"] > ctx.max_pixels:
        return False
    if xref:
        if xref not in cache:
            cache[xref] = _image_bytes(pdf, xref, pymupdf, ctx.max_pixels)
        data = cache[xref]
    else:  # an inline image: render its area of the page
        scale = CLIP_DPI / 72
        if box.width * scale * box.height * scale > ctx.max_pixels:
            return False
        data = page.get_pixmap(clip=box, dpi=CLIP_DPI).tobytes("png")
    if data is None:
        return False
    section = doc.sections[-1]
    room = section.page_width - section.left_margin - section.right_margin
    try:
        doc.add_picture(io.BytesIO(data), width=min(Pt(box.width), room))
    except Exception:  # python-docx cannot read this image format  # noqa: BLE001
        return False
    return True


def _image_bytes(pdf: Any, xref: int, pymupdf: Any, max_pixels: int) -> bytes | None:
    """A PNG or JPEG for an image object: JPEG and PNG streams as they are, the rest (JPEG 2000,
    JBIG2, CMYK, soft masks) converted to PNG. None when MuPDF cannot decode it, it decodes to
    more than ``max_pixels``, or its soft mask is not the image's size (a JPX or JBIG2 stream can
    hold more pixels than its ``/Width`` and ``/Height`` say; security review L-5)."""
    try:
        info = pdf.extract_image(xref)
        if info["ext"] in ("png", "jpeg") and not info["smask"] and info["colorspace"] in (1, 3):
            return info["image"]
        pix = pymupdf.Pixmap(pdf, xref)
        if pix.width * pix.height > max_pixels:
            return None
        if pix.alpha:
            pix = pymupdf.Pixmap(pix, 0)
        if pix.colorspace and pix.colorspace.n not in (1, 3):
            pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
        if smask := info["smask"]:
            size = (pix.width, pix.height)
            declared = tuple(int(pdf.xref_get_key(smask, k)[1]) for k in ("Width", "Height"))
            if declared != size:  # checked before decoding: a Flate mask is its declared size
                return None
            mask = pymupdf.Pixmap(pdf, smask)
            if (mask.width, mask.height) != size:
                return None
            pix = pymupdf.Pixmap(pix, mask)
        return pix.tobytes("png")
    except (RuntimeError, ValueError, pymupdf.mupdf.FzErrorBase):
        return None


def _n(count: int, word: str) -> str:
    return f"{count} {word}" if count == 1 else f"{count} {word}s"
